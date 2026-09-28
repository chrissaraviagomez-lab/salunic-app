import os
import re
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
MD_PATH = os.path.join(BASE_DIR, "Documento Explicativo del Código.md")
DOCX_PATH = os.path.join(BASE_DIR, "Documento Explicativo del Código.docx")

VERDE = RGBColor(0x0A, 0x4D, 0x2E)
AZUL = RGBColor(0x0A, 0x23, 0x42)
GRIS = RGBColor(0x44, 0x44, 0x44)
CELESTE = RGBColor(0x15, 0x65, 0xC0)

doc = Document()

doc.styles["Normal"].font.name = "Calibri"
doc.styles["Normal"].font.size = Pt(11)
doc.styles["Normal"].paragraph_format.space_after = Pt(6)


def set_run_font(run, name="Calibri"):
    run.font.name = name
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), name)
    rFonts.set(qn("w:hAnsi"), name)
    rFonts.set(qn("w:eastAsia"), name)


def shade_paragraph(paragraph, color="EFEFEF"):
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    pPr.append(shd)


def add_inline_runs(paragraph, text, base_bold=False, base_italic=False,
                    base_color=None, base_font="Calibri", base_size=None):
    tokens = re.split(r"(\*\*.*?\*\*|\*.*?\*|`.*?`)", text)
    for tok in tokens:
        if not tok:
            continue
        bold, italic, code = base_bold, base_italic, False
        content = tok
        if tok.startswith("**") and tok.endswith("**"):
            content = tok[2:-2]
            bold = True
        elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
            content = tok[1:-1]
            italic = True
        elif tok.startswith("`") and tok.endswith("`"):
            content = tok[1:-1]
            code = True
        run = paragraph.add_run(content)
        run.bold = bold
        run.italic = italic
        if code:
            set_run_font(run, "Consolas")
            run.font.color.rgb = AZUL
            run.font.size = Pt(base_size - 0.5 if base_size else 10.5)
        else:
            set_run_font(run, base_font)
            if base_color:
                run.font.color.rgb = base_color
            if base_size:
                run.font.size = Pt(base_size)


def add_code_block(lines):
    for line in lines:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(line if line else " ")
        set_run_font(run, "Consolas")
        run.font.size = Pt(9.5)
        run.font.color.rgb = GRIS
        shade_paragraph(p, "F2F5F9")


def add_image_paragraph(text, width_in=2.2):
    m = re.match(r"!\[([^\]]*)\]\(([^)]+)\)", text)
    if not m:
        return False
    alt, path = m.group(1), m.group(2)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    full = path if os.path.isabs(path) else os.path.join(BASE_DIR, path)
    if os.path.exists(full):
        run = p.add_run()
        run.add_picture(full, width=Inches(width_in))
    else:
        run = p.add_run(f"[imagen no encontrada: {path}]")
        run.font.color.rgb = GRIS
        run.italic = True
    if alt:
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = cap.add_run(alt)
        r.italic = True
        r.font.size = Pt(9)
        r.font.color.rgb = GRIS
    return True


with open(MD_PATH, "r", encoding="utf-8") as f:
    lines = f.readlines()

i = 0
in_code = False
code_buf = []
in_ul = False
in_ol = False
in_table = False
table_rows = []

while i < len(lines):
    line = lines[i].rstrip("\n")

    if line.startswith("```"):
        if in_code:
            add_code_block(code_buf)
            code_buf = []
            in_code = False
        else:
            in_code = True
        i += 1
        continue

    if in_code:
        code_buf.append(line)
        i += 1
        continue

    if not line.strip():
        in_ul = in_ol = in_table = False
        if in_table and table_rows:
            pass
        i += 1
        continue

    if line.startswith("|") and line.rstrip().endswith("|"):
        cells = [c.strip() for c in line.strip("|").split("|")]
        if re.match(r"^\s*:?-+:?\s*$", cells[0]) or all(re.match(r"^:?-+:?$", c.strip()) for c in cells):
            i += 1
            continue
        if not in_table:
            table_rows = []
            in_table = True
        table_rows.append(cells)
        if i + 1 < len(lines) and lines[i + 1].strip().startswith("|") and \
           re.match(r"^\s*\|?[:\s-]+\|?$", lines[i + 1].strip()):
            i += 2
            continue
        if i + 1 >= len(lines) or not lines[i + 1].strip().startswith("|"):
            ncols = len(table_rows[0])
            table = doc.add_table(rows=len(table_rows), cols=ncols)
            table.style = "Light Grid Accent 1"
            for r, row in enumerate(table_rows):
                for c in range(ncols):
                    cell_text = row[c] if c < len(row) else ""
                    cell_text = cell_text.replace("**", "").replace("`", "")
                    cell = table.cell(r, c)
                    cell.text = ""
                    p = cell.paragraphs[0]
                    img_m = re.match(r"!\[([^\]]*)\]\(([^)]+)\)", cell_text)
                    if img_m:
                        if r == 0:
                            add_inline_runs(p, cell_text, base_bold=True)
                        else:
                            full = img_m.group(2)
                            if not os.path.isabs(full):
                                full = os.path.join(BASE_DIR, full)
                            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                            if os.path.exists(full):
                                rn = p.add_run()
                                rn.add_picture(full, width=Inches(1.3))
                            else:
                                run = p.add_run(f"[imagen no encontrada]")
                                run.font.color.rgb = GRIS
                                run.italic = True
                    else:
                        run = p.add_run(cell_text)
                        run.font.size = Pt(9.5)
                        if r == 0:
                            run.bold = True
            in_table = False
            table_rows = []
        i += 1
        continue

    if line.startswith("# "):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(10)
        run = p.add_run(line[2:].strip())
        run.bold = True
        run.font.size = Pt(22)
        run.font.color.rgb = AZUL
        set_run_font(run)
        i += 1
        continue

    if line.startswith("## "):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(6)
        run = p.add_run(line[3:].strip())
        run.bold = True
        run.font.size = Pt(16)
        run.font.color.rgb = VERDE
        set_run_font(run)
        i += 1
        continue

    if line.startswith("### "):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        run = p.add_run(line[4:].strip())
        run.bold = True
        run.font.size = Pt(13)
        run.font.color.rgb = CELESTE
        set_run_font(run)
        i += 1
        continue

    if line.startswith("#### "):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        run = p.add_run(line[5:].strip())
        run.bold = True
        run.font.size = Pt(11.5)
        run.font.color.rgb = GRIS
        set_run_font(run)
        i += 1
        continue

    if line.startswith("> "):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.3)
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(8)
        content = line[2:].strip()
        if content.startswith("**") and content.endswith("**"):
            content = content[2:-2]
            add_inline_runs(p, content, base_bold=True, base_italic=True,
                            base_color=VERDE)
        else:
            add_inline_runs(p, content, base_italic=True, base_color=VERDE)
        shade_paragraph(p, "EAF3EE")
        i += 1
        continue

    if line.startswith("---"):
        p = doc.add_paragraph()
        pPr = p._p.get_or_add_pPr()
        pbdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "6")
        bottom.set(qn("w:space"), "1")
        bottom.set(qn("w:color"), "0A2342")
        pbdr.append(bottom)
        pPr.append(pbdr)
        i += 1
        continue

    m_ul = re.match(r"^[-*] (.*)$", line)
    if m_ul:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        add_inline_runs(p, m_ul.group(1))
        i += 1
        continue

    img_m = re.match(r"^!\[([^\]]*)\]\(([^)]+)\)", line)
    if img_m:
        add_image_paragraph(line)
        i += 1
        continue

    m_ol = re.match(r"^\d+\. (.*)$", line)
    if m_ol:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(3)
        add_inline_runs(p, m_ol.group(1))
        i += 1
        continue

    p = doc.add_paragraph()
    add_inline_runs(p, line)
    i += 1

if in_code and code_buf:
    add_code_block(code_buf)

doc.save(DOCX_PATH)
print("OK ->", DOCX_PATH)